# -*- coding: utf-8 -*-
"""
Standard Lesson Plan Builder for Learning Community (SLC) Jumping Tasks
Based on the New Taipei City standard format with #3F98DE theme.
"""
import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

PRIMARY_COLOR = "3F98DE"
LIGHT_BG = "EBF5FC"
LIGHT_GREEN_BG = "F0FFF4"
BORDER_GRAY = "CBD5E0"

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith('shd'):
            tcPr.remove(child)
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith('tcMar'):
            tcPr.remove(child)
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def format_table(tbl, col_widths, align=WD_TABLE_ALIGNMENT.CENTER):
    tbl.alignment = align
    for row in tbl.rows:
        for i, w in enumerate(col_widths):
            row.cells[i].width = Inches(w)
            set_cell_margins(row.cells[i], top=120, bottom=120, left=150, right=150)

def set_cell_text(cell, text, bold=False, font_name="微軟正黑體", font_size=10, color_rgb=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    lines = text.split("\n")
    for i, line in enumerate(lines):
        if i > 0:
            p = cell.add_paragraph()
            p.alignment = align
        r = p.add_run(line)
        r.font.name = font_name
        r.font.size = Pt(font_size)
        r.bold = bold
        if color_rgb:
            r.font.color.rgb = color_rgb

def build_lesson_plan_docx(data_dict, output_path):
    """
    Builds a complete, standard-compliant Lesson Plan DOCX file.
    """
    doc = docx.Document()
    
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)
        
    # Title Paragraphs
    p0 = doc.add_paragraph()
    r0 = p0.add_run(f"新北市立國民小學教學方案設計〔{data_dict.get('school', '學校名稱')}〕")
    r0.font.name = "標楷體"
    r0.font.size = Pt(16)
    r0.font.bold = True
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    p1 = doc.add_paragraph()
    r1 = p1.add_run(
        f"授課年級：{data_dict.get('grade', '六年級')}　　　　　　授課日期：{data_dict.get('date', '115年10月')}\n"
        f"任教學科：{data_dict.get('subject', '數學領域')}　　　　　　教 學 者（教案設計）：{data_dict.get('teacher', '教學研究團隊')}\n"
        f"單元名稱：{data_dict.get('unit_name', '第六單元 扇形面積與複合圖形探究')}　　備課成員：{data_dict.get('partners', '數學領域共備團隊')}\n"
        f"使用版本：{data_dict.get('version', '南一版（115學年度上學期，第十一冊）')}"
    )
    r1.font.name = "標楷體"
    r1.font.size = Pt(11)
    
    # Table 0: 教學設計理念
    t0 = doc.add_table(rows=2, cols=1)
    format_table(t0, [6.8])
    set_cell_background(t0.rows[0].cells[0], PRIMARY_COLOR)
    set_cell_text(t0.rows[0].cells[0], "教學設計理念", bold=True, font_size=11.5, color_rgb=RGBColor(255, 255, 255))
    set_cell_text(t0.rows[1].cells[0], data_dict.get("philosophy", ""), font_size=10)
    doc.add_paragraph()
    
    # Table 1: 核心素養
    t1 = doc.add_table(rows=4, cols=2)
    format_table(t1, [2.2, 4.6])
    t1.rows[0].cells[0].merge(t1.rows[0].cells[1])
    set_cell_background(t1.rows[0].cells[0], PRIMARY_COLOR)
    set_cell_text(t1.rows[0].cells[0], "核心素養", bold=True, font_size=11.5, color_rgb=RGBColor(255, 255, 255))
    set_cell_background(t1.rows[1].cells[0], LIGHT_BG)
    set_cell_background(t1.rows[1].cells[1], LIGHT_BG)
    set_cell_text(t1.rows[1].cells[0], "總綱", bold=True)
    set_cell_text(t1.rows[1].cells[1], "領綱具體內涵", bold=True)
    for idx, (gen, spec) in enumerate(data_dict.get("competencies", [])):
        set_cell_text(t1.rows[idx+2].cells[0], gen)
        set_cell_text(t1.rows[idx+2].cells[1], spec)
    doc.add_paragraph()
    
    # Table 2: 領域課程綱要
    t2 = doc.add_table(rows=3, cols=2)
    format_table(t2, [2.2, 4.6])
    t2.rows[0].cells[0].merge(t2.rows[0].cells[1])
    set_cell_background(t2.rows[0].cells[0], PRIMARY_COLOR)
    set_cell_text(t2.rows[0].cells[0], "領域課程綱要", bold=True, font_size=11.5, color_rgb=RGBColor(255, 255, 255))
    set_cell_background(t2.rows[1].cells[0], LIGHT_BG)
    set_cell_background(t2.rows[1].cells[1], LIGHT_BG)
    set_cell_text(t2.rows[1].cells[0], "學習表現", bold=True)
    set_cell_text(t2.rows[1].cells[1], data_dict.get("learning_performance", ""))
    set_cell_background(t2.rows[2].cells[0], LIGHT_BG)
    set_cell_text(t2.rows[2].cells[0], "學習內容與目標", bold=True)
    set_cell_text(t2.rows[2].cells[1], data_dict.get("learning_content_goals", ""))
    doc.add_paragraph()
    
    # Table 3: 教材組織與學生分析
    t3 = doc.add_table(rows=4, cols=1)
    format_table(t3, [6.8])
    set_cell_background(t3.rows[0].cells[0], PRIMARY_COLOR)
    set_cell_text(t3.rows[0].cells[0], "教材組織與學生分析", bold=True, font_size=11.5, color_rgb=RGBColor(255, 255, 255))
    set_cell_text(t3.rows[1].cells[0], "一、文本與幾何結構分析：(用圖形和文字表示)\n" + data_dict.get("structure_analysis", ""))
    set_cell_text(t3.rows[2].cells[0], "二、學生可能面臨學習的困難及發現：\n" + data_dict.get("struggles_findings", ""))
    set_cell_text(t3.rows[3].cells[0], "三、特殊學生特性描述與協同規劃：\n" + data_dict.get("special_students", ""))
    doc.add_paragraph()
    
    # Table 4: 教育四要素
    elements = data_dict.get("four_elements", [])
    t4 = doc.add_table(rows=len(elements)+1, cols=3)
    format_table(t4, [1.2, 1.8, 3.8])
    t4.rows[0].cells[0].merge(t4.rows[0].cells[2])
    set_cell_background(t4.rows[0].cells[0], PRIMARY_COLOR)
    set_cell_text(t4.rows[0].cells[0], "教學活動的教育四要素", bold=True, font_size=11.5, color_rgb=RGBColor(255, 255, 255))
    for idx, (cat, item, cont) in enumerate(elements):
        row = t4.rows[idx+1]
        set_cell_text(row.cells[0], cat, bold=True)
        set_cell_text(row.cells[1], item, bold=True)
        set_cell_text(row.cells[2], cont)
    doc.add_paragraph()
    
    # Table 5: 各節次學習重點
    t5 = doc.add_table(rows=2, cols=1)
    format_table(t5, [6.8])
    set_cell_background(t5.rows[0].cells[0], PRIMARY_COLOR)
    set_cell_text(t5.rows[0].cells[0], "各節次學習活動設計重點（本次公開課為第四節，單元統整）", bold=True, font_size=11.5, color_rgb=RGBColor(255, 255, 255))
    set_cell_text(t5.rows[1].cells[0], data_dict.get("sessions_focus", ""))
    doc.add_paragraph()
    
    # Table 6: 40分鐘教學流程
    flow_data = data_dict.get("lesson_flow", [])
    t6 = doc.add_table(rows=len(flow_data)+1, cols=4)
    format_table(t6, [1.4, 2.7, 2.1, 0.6])
    set_cell_background(t6.rows[0].cells[0], LIGHT_BG)
    set_cell_background(t6.rows[0].cells[1], LIGHT_BG)
    set_cell_background(t6.rows[0].cells[2], LIGHT_BG)
    set_cell_background(t6.rows[0].cells[3], LIGHT_BG)
    set_cell_text(t6.rows[0].cells[0], "流程", bold=True)
    set_cell_text(t6.rows[0].cells[1], "教學活動內容（含學生學習事實）", bold=True)
    set_cell_text(t6.rows[0].cells[2], "教師支援、學習引導與理答策略", bold=True)
    set_cell_text(t6.rows[0].cells[3], "時間", bold=True)
    for idx, (fl, act, sup, tm) in enumerate(flow_data):
        row = t6.rows[idx+1]
        set_cell_text(row.cells[0], fl, bold=True)
        set_cell_text(row.cells[1], act)
        set_cell_text(row.cells[2], sup)
        set_cell_text(row.cells[3], tm, bold=True)
    doc.add_paragraph()
    
    # Table 7: 理答預想表
    dialogues = data_dict.get("dialogues", [])
    t7 = doc.add_table(rows=len(dialogues)+1, cols=4)
    format_table(t7, [2.3, 1.8, 0.9, 1.8])
    set_cell_background(t7.rows[0].cells[0], LIGHT_BG)
    set_cell_background(t7.rows[0].cells[1], LIGHT_BG)
    set_cell_background(t7.rows[0].cells[2], LIGHT_BG)
    set_cell_background(t7.rows[0].cells[3], LIGHT_BG)
    set_cell_text(t7.rows[0].cells[0], "預想的學生回應/卡點", bold=True)
    set_cell_text(t7.rows[0].cells[1], "判讀與學習診斷", bold=True)
    set_cell_text(t7.rows[0].cells[2], "策略", bold=True)
    set_cell_text(t7.rows[0].cells[3], "教師具體引導台詞", bold=True)
    for idx, (resp, diag, strat, line) in enumerate(dialogues):
        row = t7.rows[idx+1]
        set_cell_text(row.cells[0], resp)
        set_cell_text(row.cells[1], diag)
        set_cell_text(row.cells[2], strat, bold=True)
        set_cell_text(row.cells[3], line)
        
    doc.save(output_path)
    print(f"Standard Lesson Plan DOCX saved at: {output_path}")
    return output_path
