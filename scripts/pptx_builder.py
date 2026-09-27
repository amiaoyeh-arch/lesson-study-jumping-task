# -*- coding: utf-8 -*-
"""
16:9 Big-Font PowerPoint Deck Builder for Learning Community (SLC) Jumping Tasks
Features:
- Standard 16:9 Widescreen (13.333 x 7.5 inches)
- Standard Header with Navy title (#1a365d) and Blue badge (#3F98DE)
- Cover Slide with left image column (no text overlap) and right text card
- Norms Slide with 50% image width cards
- Geometric Task Slides with left problem card and right clean geometry image
- Summary & Reflection Slide with quiet writing reflection illustration
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

NAVY = RGBColor(26, 54, 93)
BLUE = RGBColor(63, 152, 222) # #3F98DE
DARK_GRAY = RGBColor(45, 55, 72)
GREEN = RGBColor(39, 103, 73)
RED = RGBColor(197, 48, 48)
BG_BOX = RGBColor(247, 250, 252)
BG_BLUE = RGBColor(235, 248, 255)
BG_GREEN = RGBColor(240, 255, 244)
BORDER_COLOR = RGBColor(203, 213, 224)

def create_16x9_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs

def add_standard_header(slide, title_text, badge_text=""):
    tb = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12.133), Inches(0.85))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0
    p = tf.paragraphs[0]
    r1 = p.add_run()
    r1.text = title_text
    r1.font.name = "微軟正黑體"
    r1.font.size = Pt(25)
    r1.font.bold = True
    r1.font.color.rgb = NAVY

    if badge_text:
        r2 = p.add_run()
        r2.text = f"   [{badge_text}]"
        r2.font.name = "微軟正黑體"
        r2.font.size = Pt(16)
        r2.font.bold = True
        r2.font.color.rgb = BLUE

def add_cover_slide(prs, badge_text, title_text, subtitle_text, info_text, img_thinking_path, img_group_path):
    blank_layout = prs.slide_layouts[6]
    s = prs.slides.add_slide(blank_layout)
    bg = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(0.6), Inches(12.133), Inches(6.3))
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG_BLUE
    bg.line.color.rgb = BLUE
    bg.line.width = Pt(2)

    if os.path.exists(img_thinking_path):
        s.shapes.add_picture(img_thinking_path, Inches(1.3), Inches(1.0), Inches(2.2), Inches(2.2))
        rect_kid = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.28), Inches(0.98), Inches(2.24), Inches(2.24))
        rect_kid.fill.background()
        rect_kid.line.color.rgb = BLUE
        rect_kid.line.width = Pt(2)

    if os.path.exists(img_group_path):
        s.shapes.add_picture(img_group_path, Inches(0.9), Inches(3.6), Inches(3.0), Inches(2.25))
        rect_grp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.88), Inches(3.58), Inches(3.04), Inches(2.29))
        rect_grp.fill.background()
        rect_grp.line.color.rgb = GREEN
        rect_grp.line.width = Pt(2)

    tb = s.shapes.add_textbox(Inches(4.3), Inches(1.1), Inches(8.0), Inches(5.2))
    tf = tb.text_frame
    tf.word_wrap = True

    p_b = tf.paragraphs[0]
    r_b = p_b.add_run(f"【 {badge_text} 】\n\n")
    r_b.font.name = "微軟正黑體"
    r_b.font.size = Pt(17)
    r_b.font.bold = True
    r_b.font.color.rgb = BLUE

    p_t = tf.add_paragraph()
    r_t = p_t.add_run(f"{title_text}\n")
    r_t.font.name = "微軟正黑體"
    r_t.font.size = Pt(32)
    r_t.font.bold = True
    r_t.font.color.rgb = NAVY

    p_sub = tf.add_paragraph()
    r_sub = p_sub.add_run(f"{subtitle_text}\n\n")
    r_sub.font.name = "微軟正黑體"
    r_sub.font.size = Pt(20)
    r_sub.font.color.rgb = BLUE

    p_cap = tf.add_paragraph()
    r_cap = p_cap.add_run(info_text)
    r_cap.font.name = "微軟正黑體"
    r_cap.font.size = Pt(14)
    r_cap.font.bold = True
    r_cap.font.color.rgb = DARK_GRAY
    return s
